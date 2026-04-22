import pytest
from datetime import datetime, timezone
from pathlib import Path

from backend.automation_service.automation_service import AutomationService
from backend.automation_service.storage_service import StorageService
from backend.automation_service.excel_reporter import ExcelReporter
from backend.mail_service.models import Attachment, ParsedEmail
from backend.models.email_models import ClassificationResult, EmailCategory


def _make_email(
    uid="1",
    subject="Invoice #1",
    sender="vendor@co.com",
    body="Please pay",
    attachments=None,
) -> ParsedEmail:
    return ParsedEmail(
        uid=uid,
        subject=subject,
        sender=sender,
        date=datetime(2024, 1, 15, tzinfo=timezone.utc),
        body=body,
        attachments=attachments or [],
        raw_size=512,
    )


def _make_classification(
    category=EmailCategory.INVOICE,
    confidence=0.92,
    source="ai",
) -> ClassificationResult:
    return ClassificationResult(
        category=category,
        confidence=confidence,
        reason="billing keywords",
        source=source,
    )


def _make_attachment(filename="doc.pdf", data=b"PDF") -> Attachment:
    return Attachment(filename=filename, content_type="application/pdf", data=data)


@pytest.fixture
def service(tmp_path):
    storage = StorageService(base_dir=tmp_path)
    reporter = ExcelReporter(report_path=tmp_path / "report.xlsx")
    return AutomationService(storage=storage, reporter=reporter)


class TestProcessEmail:
    def test_returns_successful_result(self, service):
        email = _make_email()
        result = service.process(email, _make_classification())
        assert result.success is True
        assert result.email_uid == "1"
        assert result.category == "invoice"

    def test_excel_written_flag_set(self, service):
        result = service.process(_make_email(), _make_classification())
        assert result.excel_written is True

    def test_no_attachments_no_files_saved(self, service):
        result = service.process(_make_email(attachments=[]), _make_classification())
        assert result.files_saved == []
        assert result.files_skipped == []

    def test_saves_attachment_to_correct_folder(self, service, tmp_path):
        email = _make_email(attachments=[_make_attachment("invoice.pdf", b"PDF_DATA")])
        service.process(email, _make_classification(category=EmailCategory.INVOICE))
        assert (tmp_path / "invoices" / "invoice.pdf").exists()

    def test_spam_attachment_goes_to_spam_folder(self, service, tmp_path):
        email = _make_email(attachments=[_make_attachment("promo.jpg", b"IMG")])
        service.process(email, _make_classification(category=EmailCategory.SPAM))
        assert (tmp_path / "spam" / "promo.jpg").exists()

    def test_multiple_attachments_all_saved(self, service, tmp_path):
        email = _make_email(attachments=[
            _make_attachment("a.pdf", b"AAA"),
            _make_attachment("b.pdf", b"BBB"),
        ])
        result = service.process(email, _make_classification())
        assert len(result.files_saved) == 2

    def test_file_path_in_excel_row(self, service, tmp_path):
        email = _make_email(attachments=[_make_attachment("doc.pdf", b"DATA")])
        service.process(email, _make_classification())
        import openpyxl
        wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
        file_path_cell = wb.active.cell(row=2, column=6).value
        assert "doc.pdf" in file_path_cell


class TestProcessBatch:
    def test_batch_processes_all(self, service):
        pairs = [
            (_make_email(uid="1", subject="Inv 1"), _make_classification()),
            (_make_email(uid="2", subject="Inv 2"), _make_classification()),
            (_make_email(uid="3", subject="Spam", sender="x@x.com"),
             _make_classification(category=EmailCategory.SPAM)),
        ]
        results = service.process_batch(pairs)
        assert len(results) == 3
        assert all(r.success for r in results)

    def test_batch_excel_has_all_rows(self, service, tmp_path):
        pairs = [
            (_make_email(uid=str(i), sender=f"u{i}@b.com"), _make_classification())
            for i in range(4)
        ]
        service.process_batch(pairs)
        import openpyxl
        wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
        assert wb.active.max_row == 5  # 1 header + 4 data
