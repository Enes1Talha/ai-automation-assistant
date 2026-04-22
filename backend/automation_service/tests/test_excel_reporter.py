import pytest
import openpyxl
from pathlib import Path
from backend.automation_service.excel_reporter import ExcelReporter, ReportRow


def _row(
    date="2024-01-01 10:00:00",
    sender="a@b.com",
    subject="Test",
    category="invoice",
    confidence=0.9,
    file_path="/storage/invoices/doc.pdf",
    has_attachments=True,
) -> ReportRow:
    return ReportRow(
        date=date, sender=sender, subject=subject,
        category=category, confidence=confidence,
        file_path=file_path, has_attachments=has_attachments,
    )


@pytest.fixture
def reporter(tmp_path):
    return ExcelReporter(report_path=tmp_path / "report.xlsx")


class TestFileCreation:
    def test_creates_file_on_first_append(self, reporter, tmp_path):
        reporter.append(_row())
        assert (tmp_path / "report.xlsx").exists()

    def test_header_row_written(self, reporter, tmp_path):
        reporter.append(_row())
        wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
        header = [wb.active.cell(row=1, column=i).value for i in range(1, 9)]
        assert "Date" in header
        assert "Sender" in header
        assert "Category" in header

    def test_creates_parent_dirs(self, tmp_path):
        deep_path = tmp_path / "a" / "b" / "report.xlsx"
        r = ExcelReporter(report_path=deep_path)
        r.append(_row())
        assert deep_path.exists()


class TestAppend:
    def test_appends_data_row(self, reporter, tmp_path):
        reporter.append(_row(sender="x@y.com", subject="Invoice #1"))
        wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
        ws = wb.active
        assert ws.cell(row=2, column=2).value == "x@y.com"
        assert ws.cell(row=2, column=3).value == "Invoice #1"

    def test_row_count_increments(self, reporter):
        reporter.append(_row(date="2024-01-01 10:00:00", sender="a@b.com"))
        reporter.append(_row(date="2024-01-02 10:00:00", sender="c@d.com"))
        assert reporter.row_count() == 2

    def test_confidence_stored_as_float(self, reporter, tmp_path):
        reporter.append(_row(confidence=0.876))
        wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
        val = wb.active.cell(row=2, column=5).value
        assert isinstance(val, float)
        assert abs(val - 0.876) < 0.001

    def test_append_returns_true_on_new_row(self, reporter):
        assert reporter.append(_row()) is True

    def test_append_returns_false_on_duplicate(self, reporter):
        reporter.append(_row())
        assert reporter.append(_row()) is False  # same date+sender+subject

    def test_different_sender_not_duplicate(self, reporter):
        reporter.append(_row(sender="a@b.com"))
        result = reporter.append(_row(sender="c@d.com"))
        assert result is True
        assert reporter.row_count() == 2


class TestDeduplication:
    def test_exact_duplicate_not_appended(self, reporter):
        r = _row(date="2024-01-01 10:00:00", sender="x@y.com", subject="Same")
        reporter.append(r)
        reporter.append(r)
        assert reporter.row_count() == 1

    def test_different_date_not_duplicate(self, reporter):
        reporter.append(_row(date="2024-01-01 10:00:00"))
        reporter.append(_row(date="2024-01-02 10:00:00"))
        assert reporter.row_count() == 2


class TestAppendMany:
    def test_append_many_returns_stats(self, reporter):
        rows = [
            _row(date="2024-01-01 10:00:00", sender="a@b.com"),
            _row(date="2024-01-02 10:00:00", sender="c@d.com"),
            _row(date="2024-01-01 10:00:00", sender="a@b.com"),  # duplicate
        ]
        stats = reporter.append_many(rows)
        assert stats["written"] == 2
        assert stats["skipped"] == 1
        assert stats["errors"] == 0


class TestCorruptRecovery:
    def test_recreates_corrupt_file(self, tmp_path):
        path = tmp_path / "report.xlsx"
        path.write_bytes(b"NOT AN XLSX FILE")
        reporter = ExcelReporter(report_path=path)
        reporter.append(_row())
        assert reporter.row_count() == 1
