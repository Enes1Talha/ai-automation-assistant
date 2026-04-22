from __future__ import annotations

import logging
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

COLUMNS = ("date", "sender", "subject", "category", "confidence", "file_path", "has_attachments", "processed_at")

ORDER_COLUMNS = (
    "date", "sender", "subject",
    "order_number", "customer_name", "items_summary",
    "total_amount", "currency", "processed_at",
)

_HEADER_FILL = PatternFill(start_color="1A56DB", end_color="1A56DB", fill_type="solid")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_CATEGORY_COLORS = {
    "invoice": "DBEAFE",    # blue-100
    "important": "FEF3C7",  # amber-100
    "spam": "FEE2E2",       # red-100
    "other": "F3F4F6",      # gray-100
}


@dataclass
class ReportRow:
    date: str
    sender: str
    subject: str
    category: str
    confidence: float
    file_path: str
    has_attachments: bool
    processed_at: str = ""

    def __post_init__(self) -> None:
        if not self.processed_at:
            self.processed_at = datetime.now(tz=timezone.utc).isoformat(timespec="seconds")

    def as_tuple(self) -> tuple:
        return (
            self.date,
            self.sender,
            self.subject,
            self.category,
            round(self.confidence, 4),
            self.file_path,
            self.has_attachments,
            self.processed_at,
        )


@dataclass
class OrderRow:
    date: str
    sender: str
    subject: str
    order_number: str
    customer_name: str
    items_summary: str
    total_amount: float
    currency: str
    processed_at: str = ""

    def __post_init__(self) -> None:
        if not self.processed_at:
            self.processed_at = datetime.now(tz=timezone.utc).isoformat(timespec="seconds")

    def as_tuple(self) -> tuple:
        return (
            self.date,
            self.sender,
            self.subject,
            self.order_number,
            self.customer_name,
            self.items_summary,
            round(self.total_amount, 2),
            self.currency,
            self.processed_at,
        )


_ORDER_HEADER_FILL = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")


class ExcelReporter:
    """
    Appends email processing records to an Excel workbook.
    Creates the file + header row on first use.
    Deduplicates by (sender, subject, date) before appending.
    """

    def __init__(self, report_path: str | Path = "storage/report.xlsx") -> None:
        self._path = Path(report_path)
        self._path.parent.mkdir(parents=True, exist_ok=True)

    # ── Public ────────────────────────────────────────────────────────────────

    def append(self, row: ReportRow) -> bool:
        """Append a row. Returns False if duplicate, True if written."""
        wb = self._load_or_create()
        ws = wb.active

        if self._is_duplicate(ws, row):
            logger.debug("Duplicate row skipped: %s / %s", row.sender, row.subject)
            return False

        new_row = ws.max_row + 1
        for col_idx, value in enumerate(row.as_tuple(), start=1):
            cell = ws.cell(row=new_row, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=False, vertical="center")

        self._apply_row_color(ws, new_row, row.category)
        self._save(wb)
        logger.info("Appended row for %s → %s", row.sender, row.category)
        return True

    def append_order(self, row: OrderRow) -> bool:
        """Append an order row to the 'Orders' sheet. Returns False if duplicate."""
        wb = self._load_or_create()
        ws = self._get_or_create_orders_sheet(wb)

        if self._is_order_duplicate(ws, row):
            logger.debug("Duplicate order skipped: %s / %s", row.sender, row.order_number)
            return False

        new_row = ws.max_row + 1
        for col_idx, value in enumerate(row.as_tuple(), start=1):
            cell = ws.cell(row=new_row, column=col_idx, value=value)
            cell.alignment = Alignment(wrap_text=False, vertical="center")

        self._save(wb)
        logger.info("Appended order row for %s → %s", row.sender, row.order_number)
        return True

    def append_many(self, rows: list[ReportRow]) -> dict:
        written = skipped = errors = 0
        for row in rows:
            try:
                if self.append(row):
                    written += 1
                else:
                    skipped += 1
            except Exception as exc:
                logger.error("Failed to append row: %s", exc)
                errors += 1
        return {"written": written, "skipped": skipped, "errors": errors}

    def row_count(self) -> int:
        if not self._path.exists():
            return 0
        wb = openpyxl.load_workbook(self._path)
        return max(wb.active.max_row - 1, 0)  # exclude header

    # ── Private ───────────────────────────────────────────────────────────────

    def _get_or_create_orders_sheet(self, wb: openpyxl.Workbook):
        if "Orders" in wb.sheetnames:
            return wb["Orders"]
        ws = wb.create_sheet("Orders")
        self._write_orders_header(ws)
        return ws

    @staticmethod
    def _write_orders_header(ws) -> None:
        headers = [c.replace("_", " ").title() for c in ORDER_COLUMNS]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.fill = _ORDER_HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
        widths = [20, 30, 40, 18, 25, 45, 14, 10, 22]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width
        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"

    @staticmethod
    def _is_order_duplicate(ws, row: OrderRow) -> bool:
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            if (
                excel_row[0] == row.date
                and excel_row[1] == row.sender
                and excel_row[3] == row.order_number
            ):
                return True
        return False

    def _load_or_create(self) -> openpyxl.Workbook:
        if self._path.exists():
            try:
                return openpyxl.load_workbook(self._path)
            except Exception as exc:
                logger.warning("Corrupt workbook at %s — recreating: %s", self._path, exc)
                self._path.unlink()

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Email Report"
        self._write_header(ws)
        self._save(wb)
        return wb

    @staticmethod
    def _write_header(ws) -> None:
        headers = [c.replace("_", " ").title() for c in COLUMNS]
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = _HEADER_FONT
            cell.fill = _HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Column widths
        widths = [20, 30, 40, 12, 12, 50, 15, 22]
        for col_idx, width in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(col_idx)].width = width

        ws.row_dimensions[1].height = 22
        ws.freeze_panes = "A2"

    @staticmethod
    def _is_duplicate(ws, row: ReportRow) -> bool:
        # Check rows 2+ (skip header)
        for excel_row in ws.iter_rows(min_row=2, values_only=True):
            if (
                excel_row[0] == row.date
                and excel_row[1] == row.sender
                and excel_row[2] == row.subject
            ):
                return True
        return False

    @staticmethod
    def _apply_row_color(ws, row_num: int, category: str) -> None:
        hex_color = _CATEGORY_COLORS.get(category, "FFFFFF")
        fill = PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")
        for col_idx in range(1, len(COLUMNS) + 1):
            ws.cell(row=row_num, column=col_idx).fill = fill

    def _save(self, wb: openpyxl.Workbook) -> None:
        try:
            wb.save(self._path)
        except PermissionError as exc:
            raise PermissionError(f"Cannot write report (file may be open): {self._path}") from exc
