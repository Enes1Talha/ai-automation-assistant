"""
Reads order rows from the 'Orders' sheet of the Excel report.
Used by the /orders API router.
"""
from __future__ import annotations

import logging
from dataclasses import dataclass
from pathlib import Path

import openpyxl

from backend.automation_service.excel_reporter import ORDER_COLUMNS

logger = logging.getLogger(__name__)


@dataclass
class OrderRecord:
    date: str
    sender: str
    subject: str
    order_number: str
    customer_name: str
    items_summary: str
    total_amount: float
    currency: str
    processed_at: str


@dataclass
class CurrencyStat:
    currency: str
    count: int
    total: float


@dataclass
class OrderStats:
    total_orders: int
    total_revenue: float
    by_currency: list[CurrencyStat]


class OrderReader:
    """Reads structured order data from the 'Orders' Excel sheet."""

    def __init__(self, report_path: str | Path = "storage/report.xlsx") -> None:
        self._path = Path(report_path)

    def list_orders(self, page: int = 1, page_size: int = 20) -> tuple[int, list[OrderRecord]]:
        """Returns (total, page_items) sorted newest-first."""
        rows = self._read_all()
        rows.sort(key=lambda r: r.processed_at, reverse=True)
        total = len(rows)
        start = (page - 1) * page_size
        return total, rows[start: start + page_size]

    def get_stats(self) -> OrderStats:
        rows = self._read_all()
        total_revenue = sum(r.total_amount for r in rows)

        by_currency: dict[str, CurrencyStat] = {}
        for r in rows:
            key = r.currency or "N/A"
            if key not in by_currency:
                by_currency[key] = CurrencyStat(currency=key, count=0, total=0.0)
            by_currency[key].count += 1
            by_currency[key].total += r.total_amount

        return OrderStats(
            total_orders=len(rows),
            total_revenue=round(total_revenue, 2),
            by_currency=sorted(by_currency.values(), key=lambda s: s.total, reverse=True),
        )

    def _read_all(self) -> list[OrderRecord]:
        if not self._path.exists():
            return []
        try:
            wb = openpyxl.load_workbook(self._path, read_only=True, data_only=True)
        except Exception as exc:
            logger.warning("Cannot open report: %s", exc)
            return []

        if "Orders" not in wb.sheetnames:
            return []

        ws = wb["Orders"]
        records: list[OrderRecord] = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(row):
                continue
            try:
                records.append(OrderRecord(
                    date=str(row[0] or ""),
                    sender=str(row[1] or ""),
                    subject=str(row[2] or ""),
                    order_number=str(row[3] or ""),
                    customer_name=str(row[4] or ""),
                    items_summary=str(row[5] or ""),
                    total_amount=float(row[6] or 0.0),
                    currency=str(row[7] or ""),
                    processed_at=str(row[8] or ""),
                ))
            except Exception as exc:
                logger.debug("Skipping malformed order row: %s", exc)
        return records
