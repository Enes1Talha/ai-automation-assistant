"""
Tests for order extraction tool and Orders Excel sheet.
"""
import pytest
from pathlib import Path
from unittest.mock import AsyncMock, MagicMock, patch

from backend.agent.tools.executor import ToolExecutor, ToolResult
from backend.ai_service.classification_service import ClassificationService
from backend.ai_service.order_extractor import OrderData, OrderExtractorService
from backend.automation_service.excel_reporter import ExcelReporter, OrderRow
from backend.automation_service.storage_service import StorageService


def _make_executor(tmp_path: Path) -> ToolExecutor:
    classifier = MagicMock(spec=ClassificationService)
    storage = StorageService(base_dir=tmp_path)
    reporter = ExcelReporter(report_path=tmp_path / "report.xlsx")
    extractor = MagicMock(spec=OrderExtractorService)
    return ToolExecutor(
        classifier=classifier,
        storage=storage,
        reporter=reporter,
        order_extractor=extractor,
    )


# ── extract_order_data tool ───────────────────────────────────────────────────

class TestExtractOrderDataTool:
    @pytest.mark.asyncio
    async def test_returns_order_fields(self, tmp_path):
        executor = _make_executor(tmp_path)
        executor._order_extractor.extract = AsyncMock(return_value=OrderData(
            order_number="ORD-42",
            customer_name="Jane Doe",
            items_summary="Widget x2, Gadget x1",
            total_amount=150.00,
            currency="USD",
            extracted=True,
        ))
        result = await executor.execute("extract_order_data", {
            "subject": "Order Confirmation #42",
            "sender": "shop@example.com",
            "body": "Your order ORD-42 for $150.00 has been confirmed.",
        })
        assert result.success is True
        assert result.output["order_number"] == "ORD-42"
        assert result.output["customer_name"] == "Jane Doe"
        assert result.output["total_amount"] == 150.00
        assert result.output["currency"] == "USD"

    @pytest.mark.asyncio
    async def test_extraction_failure_returns_error(self, tmp_path):
        executor = _make_executor(tmp_path)
        executor._order_extractor.extract = AsyncMock(return_value=OrderData(extracted=False))
        result = await executor.execute("extract_order_data", {
            "subject": "x", "sender": "x", "body": "x"
        })
        assert result.success is False
        assert result.error is not None

    @pytest.mark.asyncio
    async def test_empty_fields_still_succeed(self, tmp_path):
        executor = _make_executor(tmp_path)
        executor._order_extractor.extract = AsyncMock(return_value=OrderData(
            order_number="",
            customer_name="",
            items_summary="",
            total_amount=0.0,
            currency="",
            extracted=True,
        ))
        result = await executor.execute("extract_order_data", {
            "subject": "Invoice", "sender": "a@b.com", "body": "pay"
        })
        assert result.success is True
        assert result.output["total_amount"] == 0.0


# ── save_order_to_excel tool ──────────────────────────────────────────────────

class TestSaveOrderToExcelTool:
    @pytest.mark.asyncio
    async def test_writes_order_row(self, tmp_path):
        executor = _make_executor(tmp_path)
        result = await executor.execute("save_order_to_excel", {
            "date": "2024-03-15 09:00:00",
            "sender": "vendor@shop.com",
            "subject": "Order #99",
            "order_number": "ORD-99",
            "customer_name": "Alice",
            "items_summary": "Widget x5",
            "total_amount": 250.0,
            "currency": "EUR",
        })
        assert result.success is True
        assert result.output["written"] is True

    @pytest.mark.asyncio
    async def test_duplicate_order_skipped(self, tmp_path):
        executor = _make_executor(tmp_path)
        inp = {
            "date": "2024-03-15 09:00:00",
            "sender": "vendor@shop.com",
            "subject": "Order #99",
            "order_number": "ORD-99",
            "customer_name": "Alice",
            "items_summary": "Widget x5",
            "total_amount": 250.0,
            "currency": "EUR",
        }
        await executor.execute("save_order_to_excel", inp)
        result = await executor.execute("save_order_to_excel", inp)
        assert result.output["skipped_duplicate"] is True

    @pytest.mark.asyncio
    async def test_missing_optional_fields_use_defaults(self, tmp_path):
        executor = _make_executor(tmp_path)
        result = await executor.execute("save_order_to_excel", {
            "date": "2024-03-15 09:00:00",
            "sender": "x@x.com",
            "subject": "Invoice",
            "order_number": "INV-001",
            "total_amount": 99.9,
        })
        assert result.success is True


# ── ExcelReporter Orders sheet ────────────────────────────────────────────────

class TestExcelReporterOrdersSheet:
    def test_orders_sheet_created_on_first_append(self, tmp_path):
        import openpyxl
        reporter = ExcelReporter(report_path=tmp_path / "r.xlsx")
        row = OrderRow(
            date="2024-01-01 10:00:00",
            sender="a@b.com",
            subject="Invoice #1",
            order_number="ORD-1",
            customer_name="Bob",
            items_summary="Item x1",
            total_amount=50.0,
            currency="USD",
        )
        written = reporter.append_order(row)
        assert written is True
        wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
        assert "Orders" in wb.sheetnames

    def test_orders_sheet_header_has_correct_columns(self, tmp_path):
        import openpyxl
        reporter = ExcelReporter(report_path=tmp_path / "r.xlsx")
        reporter.append_order(OrderRow(
            date="2024-01-01 10:00:00", sender="a@b.com", subject="Inv",
            order_number="X", customer_name="Y", items_summary="Z",
            total_amount=1.0, currency="USD",
        ))
        wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
        ws = wb["Orders"]
        headers = [ws.cell(row=1, column=i).value for i in range(1, 10)]
        assert "Order Number" in headers
        assert "Total Amount" in headers
        assert "Currency" in headers

    def test_duplicate_order_not_appended(self, tmp_path):
        reporter = ExcelReporter(report_path=tmp_path / "r.xlsx")
        row = OrderRow(
            date="2024-01-01 10:00:00", sender="a@b.com", subject="Inv",
            order_number="ORD-DUP", customer_name="C", items_summary="I",
            total_amount=10.0, currency="USD",
        )
        assert reporter.append_order(row) is True
        assert reporter.append_order(row) is False

    def test_email_sheet_unaffected_by_order_append(self, tmp_path):
        import openpyxl
        from backend.automation_service.excel_reporter import ReportRow
        reporter = ExcelReporter(report_path=tmp_path / "r.xlsx")
        reporter.append(ReportRow(
            date="2024-01-01", sender="x@x.com", subject="Hi",
            category="invoice", confidence=0.9, file_path="", has_attachments=False,
        ))
        reporter.append_order(OrderRow(
            date="2024-01-01", sender="x@x.com", subject="Hi",
            order_number="ORD-1", customer_name="X", items_summary="",
            total_amount=0.0, currency="",
        ))
        wb = openpyxl.load_workbook(tmp_path / "r.xlsx")
        assert wb.active.title == "Email Report"
        assert "Orders" in wb.sheetnames


# ── OrderExtractorService unit tests ─────────────────────────────────────────

class TestOrderExtractorService:
    @pytest.mark.asyncio
    async def test_parses_valid_json_response(self):
        service = OrderExtractorService(api_key="test")
        mock_block = MagicMock()
        mock_block.text = '{"order_number":"ORD-7","customer_name":"Eve","items_summary":"A x1","total_amount":99.0,"currency":"GBP"}'
        mock_response = MagicMock()
        mock_response.content = [mock_block]
        service._client = MagicMock()
        service._client.messages.create.return_value = mock_response

        result = await service.extract("Inv", "s@s.com", "body")
        assert result.extracted is True
        assert result.order_number == "ORD-7"
        assert result.currency == "GBP"
        assert result.total_amount == 99.0

    @pytest.mark.asyncio
    async def test_handles_markdown_fenced_json(self):
        service = OrderExtractorService(api_key="test")
        mock_block = MagicMock()
        mock_block.text = '```json\n{"order_number":"X","customer_name":"","items_summary":"","total_amount":0,"currency":""}\n```'
        mock_response = MagicMock()
        mock_response.content = [mock_block]
        service._client = MagicMock()
        service._client.messages.create.return_value = mock_response

        result = await service.extract("s", "e", "b")
        assert result.extracted is True
        assert result.order_number == "X"

    @pytest.mark.asyncio
    async def test_api_failure_returns_not_extracted(self):
        service = OrderExtractorService(api_key="test")
        service._client = MagicMock()
        service._client.messages.create.side_effect = Exception("API error")

        result = await service.extract("s", "e", "b")
        assert result.extracted is False

    @pytest.mark.asyncio
    async def test_invalid_json_returns_not_extracted(self):
        service = OrderExtractorService(api_key="test")
        mock_block = MagicMock()
        mock_block.text = "this is not json at all"
        mock_response = MagicMock()
        mock_response.content = [mock_block]
        service._client = MagicMock()
        service._client.messages.create.return_value = mock_response

        result = await service.extract("s", "e", "b")
        assert result.extracted is False
